from __future__ import annotations

import hashlib
import hmac
import io
import json
import math
import os
import re
import secrets
import sqlite3
import time
import uuid
import asyncio
import xml.etree.ElementTree as ET
from collections import Counter, defaultdict
from contextlib import asynccontextmanager
from datetime import datetime, timedelta, timezone
from pathlib import Path
from tempfile import gettempdir
from typing import Any, AsyncIterator

import httpx
from fastapi import FastAPI, Header, HTTPException, Request, WebSocket, WebSocketDisconnect
from fastapi.responses import Response, StreamingResponse
from jinja2 import Environment
from jose import jwt
from pydantic import BaseModel, EmailStr, Field
from docx import Document

JWT_SECRET = os.getenv("JWT_SECRET", "dev-secret-change-me")
JWT_ALGORITHM = "HS256"
ACCESS_TTL_SECONDS = 3600
OLLAMA_BASE_URL = os.getenv("OLLAMA_BASE_URL", "http://localhost:11434").rstrip("/")
OLLAMA_MODEL = os.getenv("OLLAMA_MODEL", "gemma4:e2b")
OLLAMA_KEEP_ALIVE = os.getenv("OLLAMA_KEEP_ALIVE", "-1")
SUPERDOC_MCP_URL = os.getenv("SUPERDOC_MCP_URL", "http://localhost:8090/mcp")
MAX_CONTEXT_MESSAGES = int(os.getenv("MAX_CONTEXT_MESSAGES", "12"))
CHAT_DB_PATH = Path(os.getenv("CHAT_DB_PATH", str(Path(gettempdir()) / "scholarflow" / "chat_history.db")))
PALETTE_DB_PATH = Path(
    os.getenv("PALETTE_DB_PATH", str(Path(gettempdir()) / "scholarflow" / "color_palettes.json"))
)
DOCX_ARTIFACT_DIR = Path(os.getenv("DOCX_ARTIFACT_DIR", str(Path(gettempdir()) / "scholarflow" / "docx_artifacts")))
DEFAULT_CONVERSATION_TITLE = "New chat"
MAX_CONVERSATION_TITLE_LENGTH = 60

users_by_google_sub: dict[str, dict[str, Any]] = {}
users_by_email: dict[str, dict[str, Any]] = {}
refresh_tokens_by_user: dict[str, str] = {}
service_registry: dict[str, dict[str, Any]] = {}
subscriptions: dict[str, str] = defaultdict(lambda: "free")
app_started_at = time.time()
app_instance_id = os.getenv("SERVICE_INSTANCE_ID", f"core-{uuid.uuid4()}")
app_service_name = os.getenv("SERVICE_NAME", "core")
app_metadata_tags = [tag for tag in os.getenv("SERVICE_METADATA_TAGS", "core,fastapi,scholarflow").split(",") if tag]
plagiarism_events: list[dict[str, Any]] = []
doc_crdt_state: dict[str, dict[str, Any]] = defaultdict(
    lambda: {"clock": 0, "text": "", "cursor": {}, "operations": [], "applied_operation_ids": set()}
)
chat_db_initialized = False
palette_db_initialized = False
DOCX_ARTIFACT_TTL_SECONDS = 1800
docx_artifacts: dict[str, dict[str, Any]] = {}
superdoc_mcp_session_id: str | None = None
superdoc_tool_name_cache: dict[str, str] = {}
CHAT_AGENT_PROFILES: list[dict[str, str]] = [
    {
        "name": "Analyst",
        "system_prompt": (
            "You are Analyst. Provide a complete, final answer focused on structure, constraints, and a clear plan. "
            "Do not reference other agents."
        ),
    },
    {
        "name": "Researcher",
        "system_prompt": (
            "You are Researcher. Provide a complete, final answer focused on evidence, assumptions, and risks. "
            "Do not reference other agents."
        ),
    },
    {
        "name": "Strategist",
        "system_prompt": (
            "You are Strategist. Provide a complete, final answer focused on decisions, tradeoffs, and execution steps. "
            "Do not reference other agents."
        ),
    },
]


def _utc_now_iso() -> str:
    return datetime.now(timezone.utc).isoformat()


def _connect_chat_db() -> sqlite3.Connection:
    _init_chat_db()
    return sqlite3.connect(CHAT_DB_PATH)


def _init_palette_db() -> None:
    global palette_db_initialized
    if palette_db_initialized:
        return
    PALETTE_DB_PATH.parent.mkdir(parents=True, exist_ok=True)
    if not PALETTE_DB_PATH.exists():
        PALETTE_DB_PATH.write_text("[]", encoding="utf-8")
    palette_db_initialized = True


def _read_palette_docs() -> list[dict[str, Any]]:
    _init_palette_db()
    try:
        payload = json.loads(PALETTE_DB_PATH.read_text(encoding="utf-8") or "[]")
    except json.JSONDecodeError:
        payload = []
    if not isinstance(payload, list):
        return []
    return [entry for entry in payload if isinstance(entry, dict)]


def _write_palette_docs(items: list[dict[str, Any]]) -> None:
    _init_palette_db()
    safe_items = [entry for entry in items if isinstance(entry, dict)]
    temp_path = PALETTE_DB_PATH.with_suffix(".tmp")
    temp_path.write_text(json.dumps(safe_items, ensure_ascii=False, indent=2), encoding="utf-8")
    temp_path.replace(PALETTE_DB_PATH)


def _list_color_palettes() -> list[dict[str, Any]]:
    docs = _read_palette_docs()
    docs.sort(key=lambda item: str(item.get("updated_at") or item.get("created_at") or ""), reverse=True)
    return docs


def _create_color_palette(name: str, colors: dict[str, Any]) -> dict[str, Any]:
    docs = _read_palette_docs()
    now = _utc_now_iso()
    entry = {
        "id": str(uuid.uuid4()),
        "name": name.strip(),
        "colors": colors,
        "created_at": now,
        "updated_at": now,
    }
    docs.append(entry)
    _write_palette_docs(docs)
    return entry


def _delete_color_palette(palette_id: str) -> bool:
    docs = _read_palette_docs()
    filtered = [entry for entry in docs if str(entry.get("id")) != palette_id]
    if len(filtered) == len(docs):
        return False
    _write_palette_docs(filtered)
    return True


def _init_chat_db() -> None:
    global chat_db_initialized
    if chat_db_initialized:
        return
    CHAT_DB_PATH.parent.mkdir(parents=True, exist_ok=True)
    with sqlite3.connect(CHAT_DB_PATH) as conn:
        conn.execute(
            """
            CREATE TABLE IF NOT EXISTS conversations (
                id TEXT PRIMARY KEY,
                title TEXT NOT NULL,
                created_at TEXT NOT NULL,
                updated_at TEXT NOT NULL
            )
            """
        )
        conn.execute(
            """
            CREATE TABLE IF NOT EXISTS messages (
                id TEXT PRIMARY KEY,
                conversation_id TEXT NOT NULL,
                role TEXT NOT NULL,
                content TEXT NOT NULL,
                created_at TEXT NOT NULL,
                FOREIGN KEY(conversation_id) REFERENCES conversations(id) ON DELETE CASCADE
            )
            """
        )
        conn.execute("CREATE INDEX IF NOT EXISTS idx_messages_conversation ON messages(conversation_id, created_at)")
        conn.execute("CREATE INDEX IF NOT EXISTS idx_conversations_updated ON conversations(updated_at DESC)")
    chat_db_initialized = True


def _conversation_exists(conversation_id: str) -> bool:
    with _connect_chat_db() as conn:
        row = conn.execute("SELECT 1 FROM conversations WHERE id = ?", (conversation_id,)).fetchone()
    return bool(row)


def _create_conversation(title: str = DEFAULT_CONVERSATION_TITLE) -> dict[str, Any]:
    conversation_id = str(uuid.uuid4())
    now = _utc_now_iso()
    with _connect_chat_db() as conn:
        conn.execute(
            "INSERT INTO conversations (id, title, created_at, updated_at) VALUES (?, ?, ?, ?)",
            (conversation_id, title, now, now),
        )
    return {"id": conversation_id, "title": title, "message_count": 0, "updated_at": now}


def _truncate_title(value: str) -> str:
    clean = " ".join(value.strip().split())
    if len(clean) <= MAX_CONVERSATION_TITLE_LENGTH:
        return clean
    if MAX_CONVERSATION_TITLE_LENGTH <= 3:
        return clean[:MAX_CONVERSATION_TITLE_LENGTH]
    sliced = clean[: MAX_CONVERSATION_TITLE_LENGTH - 3].rstrip()
    if " " in sliced:
        sliced = sliced.rsplit(" ", 1)[0]
    return f"{sliced}..."


def _add_message(conversation_id: str, role: str, content: str) -> dict[str, Any]:
    message_id = str(uuid.uuid4())
    now = _utc_now_iso()
    with _connect_chat_db() as conn:
        conn.execute(
            "INSERT INTO messages (id, conversation_id, role, content, created_at) VALUES (?, ?, ?, ?, ?)",
            (message_id, conversation_id, role, content, now),
        )
        conn.execute("UPDATE conversations SET updated_at = ? WHERE id = ?", (now, conversation_id))
        row = conn.execute("SELECT title FROM conversations WHERE id = ?", (conversation_id,)).fetchone()
        if row and row[0].strip() == DEFAULT_CONVERSATION_TITLE and role == "user":
            conn.execute(
                "UPDATE conversations SET title = ? WHERE id = ?",
                (_truncate_title(content), conversation_id),
            )
    return {"id": message_id, "conversation_id": conversation_id, "role": role, "content": content, "created_at": now}


def _list_conversations() -> list[dict[str, Any]]:
    with _connect_chat_db() as conn:
        rows = conn.execute(
            """
            SELECT c.id, c.title, c.updated_at, COUNT(m.id) AS message_count
            FROM conversations c
            LEFT JOIN messages m ON m.conversation_id = c.id
            GROUP BY c.id
            ORDER BY c.updated_at DESC
            """
        ).fetchall()
    return [
        {
            "id": row[0],
            "title": row[1],
            "updated_at": row[2],
            "message_count": row[3],
        }
        for row in rows
    ]


def _list_messages(conversation_id: str) -> list[dict[str, Any]]:
    with _connect_chat_db() as conn:
        rows = conn.execute(
            """
            SELECT id, role, content, created_at
            FROM messages
            WHERE conversation_id = ?
            ORDER BY created_at ASC
            """,
            (conversation_id,),
        ).fetchall()
    return [{"id": row[0], "role": row[1], "content": row[2], "created_at": row[3]} for row in rows]


def register_self() -> None:
    service_registry[app_instance_id] = {
        "instance_id": app_instance_id,
        "service_name": app_service_name,
        "metadata_tags": app_metadata_tags,
        "uptime_seconds": 0,
        "request_latency_ms": 12,
        "error_rate": 0.0,
    }


def _normalized_keep_alive() -> str:
    raw = OLLAMA_KEEP_ALIVE.strip()
    if raw == "-1":
        # Ollama expects duration strings on some endpoints; this value means effectively no eviction.
        return "2562047h47m16.854775807s"
    if re.fullmatch(r"-?\d+", raw):
        return f"{raw}s"
    return raw


async def _warm_ollama_model() -> None:
    payload = {
        "model": OLLAMA_MODEL,
        "stream": False,
        "keep_alive": _normalized_keep_alive(),
        "prompt": "Warmup",
        "options": {"num_predict": 1},
    }
    timeout = httpx.Timeout(timeout=30.0, connect=5.0)
    try:
        async with httpx.AsyncClient(timeout=timeout) as http_client:
            await http_client.post(f"{OLLAMA_BASE_URL}/api/generate", json=payload)
    except httpx.HTTPError:
        # Warmup is best-effort and should never break API startup.
        return


@asynccontextmanager
async def lifespan(_: FastAPI):
    register_self()
    asyncio.create_task(_warm_ollama_model())
    yield


app = FastAPI(title="ScholarFlow AI Enterprise", lifespan=lifespan)


class RegisterServiceRequest(BaseModel):
    service_name: str
    metadata_tags: list[str] = Field(default_factory=list)


class GoogleProfile(BaseModel):
    sub: str
    email: EmailStr
    name: str


class OAuthCallbackRequest(BaseModel):
    code: str
    profile: GoogleProfile
    tier: str = Field(default="free", pattern="^(free|enterprise)$")


class EmailPasswordRegisterRequest(BaseModel):
    email: EmailStr
    name: str
    password: str = Field(min_length=8)
    tier: str = Field(default="free", pattern="^(free|enterprise)$")


class EmailPasswordLoginRequest(BaseModel):
    email: EmailStr
    password: str


class ChatMessage(BaseModel):
    role: str = Field(pattern="^(system|user|assistant)$")
    content: str = Field(min_length=1)


class ChatStreamRequest(BaseModel):
    messages: list[ChatMessage] = Field(min_length=1)
    conversation_id: str | None = None


class ConversationCreateRequest(BaseModel):
    title: str | None = Field(default=None, max_length=MAX_CONVERSATION_TITLE_LENGTH)


class PaletteColorEntry(BaseModel):
    color: str = Field(pattern=r"^#[0-9a-fA-F]{6}$")
    isLocked: bool = False


class PaletteColorsPayload(BaseModel):
    textColor: PaletteColorEntry
    backgroundColor: PaletteColorEntry
    primaryColor: PaletteColorEntry
    secondaryColor: PaletteColorEntry
    accentColor: PaletteColorEntry


class ColorPaletteCreateRequest(BaseModel):
    name: str = Field(min_length=1, max_length=80)
    colors: PaletteColorsPayload


class SimilarityRequest(BaseModel):
    text: str
    sources: list[str]


class RemediationRequest(BaseModel):
    flagged_sentences: list[str]


class CrdtOperation(BaseModel):
    op_id: str
    actor_id: str
    timestamp: int
    text_patch: str = ""
    cursor_position: int = 0


class CitationClaim(BaseModel):
    claim: str
    source_text: str


class CitationValidationRequest(BaseModel):
    citations: list[CitationClaim]


class ExportRequest(BaseModel):
    title: str
    author: str
    body: str
    custom_css: str = ""


class DocxTemplateRequest(BaseModel):
    template_xml: str = Field(min_length=1)
    json_data: dict[str, Any] = Field(default_factory=dict)
    filename: str = Field(default="draft.docx", max_length=120)


class DocxPreviewResponse(BaseModel):
    resolved_xml: str
    preview_html: str
    paragraphs: list[str]


class DocxSessionResponse(BaseModel):
    docx_id: str
    filename: str
    viewer_path: str


class DocxToolGenerateRequest(BaseModel):
    conversation_id: str | None = None
    prompt: str | None = None
    messages: list[ChatMessage] = Field(default_factory=list)
    current_template_xml: str | None = None
    current_json_data: dict[str, Any] | None = None


class DocxToolGenerateResponse(BaseModel):
    template_xml: str
    json_data: dict[str, Any]
    filename: str
    model: str
    source: str


class SuperdocOpenRequest(BaseModel):
    docx_id: str = Field(min_length=1)


class SuperdocApplyRequest(BaseModel):
    docx_id: str = Field(min_length=1)
    content: str = Field(min_length=1)
    suggest: bool = False


class SuperdocSessionRequest(BaseModel):
    session_id: str = Field(min_length=1)


def _tokenize(text: str) -> Counter:
    tokens = re.findall(r"[a-zA-Z0-9]+", text.lower())
    return Counter(tokens)


def _cosine_similarity(a: str, b: str) -> float:
    va = _tokenize(a)
    vb = _tokenize(b)
    if not va or not vb:
        return 0.0
    dot = sum(va[t] * vb.get(t, 0) for t in va)
    norm_a = math.sqrt(sum(v * v for v in va.values()))
    norm_b = math.sqrt(sum(v * v for v in vb.values()))
    return dot / (norm_a * norm_b) if norm_a and norm_b else 0.0


def _split_sentences(text: str) -> list[str]:
    return [s.strip() for s in re.split(r"(?<=[.!?])\s+", text.strip()) if s.strip()]


def _trace_id_from_request(headers: dict[str, str]) -> str:
    traceparent = headers.get("traceparent", "")
    if traceparent:
        parts = traceparent.split("-")
        if len(parts) >= 2 and parts[1]:
            return parts[1]
    return headers.get("x-request-id", str(uuid.uuid4()))


def _lookup_path(data: dict[str, Any], dotted_path: str) -> str:
    current: Any = data
    for key in dotted_path.split("."):
        if isinstance(current, dict) and key in current:
            current = current[key]
            continue
        return ""
    if isinstance(current, (dict, list)):
        return json.dumps(current)
    return "" if current is None else str(current)


def _resolve_template_placeholders(template_xml: str, json_data: dict[str, Any]) -> str:
    def replace_match(match: re.Match[str]) -> str:
        key = match.group(1).strip()
        return _lookup_path(json_data, key)

    return re.sub(r"\{\{\s*([a-zA-Z0-9_.-]+)\s*\}\}", replace_match, template_xml)


def _extract_docx_paragraphs(root: ET.Element) -> list[str]:
    paragraphs: list[str] = []
    for node in root.iter():
        if node.tag in {"p", "h1", "h2", "h3", "li"}:
            content = "".join(node.itertext()).strip()
            if content:
                paragraphs.append(content)
    if paragraphs:
        return paragraphs
    flattened = "".join(root.itertext()).strip()
    return [flattened] if flattened else []


def _build_preview_html(root: ET.Element) -> str:
    lines: list[str] = []
    for node in root:
        if node.tag in {"h1", "h2", "h3", "p"}:
            lines.append(f"<{node.tag}>{''.join(node.itertext()).strip()}</{node.tag}>")
        elif node.tag == "ul":
            items = [f"<li>{''.join(item.itertext()).strip()}</li>" for item in node.findall("li")]
            lines.append(f"<ul>{''.join(items)}</ul>")
    if not lines:
        text = "".join(root.itertext()).strip()
        if text:
            lines.append(f"<p>{text}</p>")
    return "".join(lines)


def _build_docx_bytes_from_xml(resolved_xml: str) -> bytes:
    try:
        root = ET.fromstring(resolved_xml)
    except ET.ParseError as exc:
        raise HTTPException(status_code=400, detail=f"invalid template XML: {exc}") from exc

    document = Document()
    paragraphs = _extract_docx_paragraphs(root)
    if not paragraphs:
        raise HTTPException(status_code=400, detail="template produced empty content")

    for paragraph in paragraphs:
        document.add_paragraph(paragraph)

    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def _clean_docx_artifacts() -> None:
    cutoff = time.time() - DOCX_ARTIFACT_TTL_SECONDS
    stale_ids = [artifact_id for artifact_id, entry in docx_artifacts.items() if entry["created_at"] < cutoff]
    for artifact_id in stale_ids:
        artifact = docx_artifacts.pop(artifact_id, None)
        if artifact:
            file_path = artifact.get("file_path")
            if isinstance(file_path, str) and file_path:
                Path(file_path).unlink(missing_ok=True)


def _ensure_docx_artifact_file(docx_id: str) -> str:
    _clean_docx_artifacts()
    artifact = docx_artifacts.get(docx_id)
    if not artifact:
        raise HTTPException(status_code=404, detail="docx artifact not found")

    DOCX_ARTIFACT_DIR.mkdir(parents=True, exist_ok=True)
    file_path = artifact.get("file_path")
    if isinstance(file_path, str) and file_path:
        path = Path(file_path)
    else:
        filename = artifact.get("filename") or f"{docx_id}.docx"
        safe_name = re.sub(r"[^a-zA-Z0-9._-]+", "-", str(filename)) or f"{docx_id}.docx"
        path = DOCX_ARTIFACT_DIR / f"{docx_id}-{safe_name}"
        artifact["file_path"] = str(path)

    if not path.exists():
        content = artifact.get("content")
        if not isinstance(content, (bytes, bytearray)):
            raise HTTPException(status_code=500, detail="docx artifact content unavailable")
        path.write_bytes(bytes(content))
    return str(path)


def _reload_docx_artifact_content(docx_id: str) -> None:
    artifact = docx_artifacts.get(docx_id)
    if not artifact:
        return
    file_path = artifact.get("file_path")
    if isinstance(file_path, str) and file_path and Path(file_path).exists():
        artifact["content"] = Path(file_path).read_bytes()


def _parse_streamable_http_body(raw_text: str) -> dict[str, Any]:
    body = raw_text.strip()
    if not body:
        return {}

    if body.startswith("event:") or body.startswith("data:"):
        for line in body.splitlines():
            if line.startswith("data:"):
                data = line[len("data:") :].strip()
                if not data:
                    continue
                try:
                    return json.loads(data)
                except json.JSONDecodeError:
                    continue
        return {}

    try:
        return json.loads(body)
    except json.JSONDecodeError as exc:
        raise HTTPException(status_code=502, detail=f"invalid mcp response body: {exc}") from exc


def _coerce_crew_output_text(output: Any) -> str:
    if isinstance(output, str):
        return output
    for attr in ("raw", "output", "result", "final_output"):
        if hasattr(output, attr):
            value = getattr(output, attr)
            if isinstance(value, str) and value.strip():
                return value
    return str(output)


def _extract_superdoc_session_id(parsed_payload: Any) -> str | None:
    if isinstance(parsed_payload, dict):
        direct_keys = ("session_id", "sessionId", "id")
        for key in direct_keys:
            value = parsed_payload.get(key)
            if isinstance(value, str) and value.strip():
                return value.strip()

        for nested_key in ("data", "result", "session"):
            nested_value = parsed_payload.get(nested_key)
            resolved = _extract_superdoc_session_id(nested_value)
            if resolved:
                return resolved

    if isinstance(parsed_payload, str):
        match = re.search(r'"session[_ ]?id"\s*:\s*"([^"]+)"', parsed_payload, re.IGNORECASE)
        if match:
            return match.group(1)
    return None


def _resolve_superdoc_tool_name(action: str) -> str:
    cached = superdoc_tool_name_cache.get(action)
    if cached:
        return cached

    listing = _superdoc_mcp_rpc("tools/list", {}, include_session=True)
    tools = listing.get("result", {}).get("tools", [])
    if not isinstance(tools, list):
        raise HTTPException(status_code=502, detail="superdoc mcp tools/list returned invalid payload")

    candidates: list[str] = []
    for item in tools:
        if not isinstance(item, dict):
            continue
        name = str(item.get("name") or "").strip()
        if not name:
            continue
        lowered = name.lower()
        if action in lowered:
            candidates.append(name)

    scored: list[tuple[int, str]] = []
    for name in candidates:
        lowered = name.lower()
        score = 0
        if lowered == f"superdoc_{action}":
            score += 100
        if lowered.endswith(f".{action}") or lowered.endswith(f"_{action}"):
            score += 60
        if "superdoc" in lowered:
            score += 30
        if lowered == action:
            score += 20
        scored.append((score, name))

    if not scored:
        available = [str(item.get("name")) for item in tools if isinstance(item, dict) and item.get("name")]
        raise HTTPException(
            status_code=502,
            detail=f"superdoc mcp tool for '{action}' not found. available={available}",
        )

    scored.sort(key=lambda item: item[0], reverse=True)
    resolved_name = scored[0][1]
    superdoc_tool_name_cache[action] = resolved_name
    return resolved_name


def _run_docx_generation_crew(
    seed_text: str,
    source_messages: list[dict[str, str]],
    current_template_xml: str | None = None,
    current_json_data: dict[str, Any] | None = None,
) -> tuple[str, dict[str, Any], str]:
    try:
        from crewai import Agent, Crew, LLM, Process, Task
    except Exception as exc:  # pragma: no cover - dependency/runtime validation
        raise HTTPException(status_code=500, detail=f"crewai dependency unavailable: {exc}") from exc

    transcript_lines = [f"{item['role']}: {item['content']}" for item in source_messages if item.get("content")]
    transcript = "\n".join(transcript_lines[-20:])
    current_json_text = json.dumps(current_json_data or {}, ensure_ascii=True)

    llm = LLM(
        model=os.getenv("CREWAI_MODEL", f"ollama/{OLLAMA_MODEL}"),
        base_url=os.getenv("CREWAI_BASE_URL", OLLAMA_BASE_URL),
    )

    planner = Agent(
        role="Academic Planning Strategist",
        goal="Design a rigorous document plan that addresses user intent and context gaps.",
        backstory="Expert in structuring scholarly documents and identifying missing evidence.",
        llm=llm,
        verbose=False,
        allow_delegation=False,
    )
    researcher = Agent(
        role="Evidence Synthesizer",
        goal="Extract the strongest factual anchors and unresolved questions from transcript context.",
        backstory="Specialist in balancing confidence with explicit uncertainty for research writing.",
        llm=llm,
        verbose=False,
        allow_delegation=False,
    )
    writer = Agent(
        role="DOCX Draft Composer",
        goal="Convert planning and evidence notes into clear prose sections suitable for DOCX output.",
        backstory="Professional technical writer focused on clarity, traceability, and concise structure.",
        llm=llm,
        verbose=False,
        allow_delegation=False,
    )
    editor = Agent(
        role="JSON Contract Editor",
        goal="Return only valid JSON matching the required schema for the DOCX generation endpoint.",
        backstory="Schema-first editor who enforces strict response formats without commentary.",
        llm=llm,
        verbose=False,
        allow_delegation=False,
    )

    shared_context = (
        f"Seed focus: {seed_text or 'general summary'}\n"
        f"Current template xml:\n{current_template_xml or '(none)'}\n"
        f"Current json data:\n{current_json_text}\n"
        f"Transcript:\n{transcript}"
    )

    planning_task = Task(
        description=(
            "Analyze user intent and produce an exhaustive section plan with priority ordering, "
            "scope boundaries, and any assumptions that must be made.\n\n"
            f"{shared_context}"
        ),
        expected_output="Bullet plan with assumptions and recommended sections.",
        agent=planner,
    )
    research_task = Task(
        description=(
            "From the transcript, enumerate evidence anchors, contradictions, and unknowns to address in the draft. "
            "Be explicit where confidence is low."
        ),
        expected_output="Evidence synthesis with confidence notes.",
        agent=researcher,
        context=[planning_task],
    )
    writing_task = Task(
        description=(
            "Draft a polished document body using the plan and synthesis. Keep tone professional and actionable."
        ),
        expected_output="Draft content with title, abstract, and section prose.",
        agent=writer,
        context=[planning_task, research_task],
    )
    final_json_task = Task(
        description=(
            "Return ONLY JSON with keys: title, author (object with name), abstract, "
            "sections (array of objects with heading/content), filename. No markdown. No explanation."
        ),
        expected_output="Valid JSON object only.",
        agent=editor,
        context=[planning_task, research_task, writing_task],
    )

    crew = Crew(
        agents=[planner, researcher, writer, editor],
        tasks=[planning_task, research_task, writing_task, final_json_task],
        process=Process.sequential,
        verbose=False,
    )

    kickoff_result = crew.kickoff()
    output_text = _coerce_crew_output_text(kickoff_result).strip()
    parsed = _extract_json_object(output_text)
    if not parsed:
        if current_template_xml and current_json_data:
            filename_seed = _truncate_title(seed_text or str(current_json_data.get("title") or "document"))
            safe_filename = f"{re.sub(r'[^a-z0-9]+', '-', filename_seed.lower()).strip('-') or 'document'}.docx"
            return current_template_xml, current_json_data, safe_filename
        return _default_docx_draft(seed_text)

    return _normalize_docx_plan(parsed, seed_text)


async def _generate_independent_agent_reply(
    agent_name: str,
    system_prompt: str,
    context_messages: list[dict[str, str]],
) -> dict[str, str]:
    payload = {
        "model": OLLAMA_MODEL,
        "stream": False,
        "keep_alive": _normalized_keep_alive(),
        "messages": [{"role": "system", "content": system_prompt}] + context_messages,
    }

    timeout = httpx.Timeout(timeout=300.0, connect=10.0)
    async with httpx.AsyncClient(timeout=timeout) as http_client:
        response = await http_client.post(f"{OLLAMA_BASE_URL}/api/chat", json=payload)
        response.raise_for_status()
        response_payload = response.json()

    content = str(response_payload.get("message", {}).get("content", "")).strip()
    if not content:
        raise HTTPException(status_code=502, detail=f"agent '{agent_name}' returned empty output")
    return {"agent": agent_name, "content": content}


async def _generate_chatdev_like_responses(context_messages: list[dict[str, str]]) -> list[dict[str, str]]:
    tasks = [
        asyncio.create_task(_generate_independent_agent_reply(profile["name"], profile["system_prompt"], context_messages))
        for profile in CHAT_AGENT_PROFILES
    ]

    results: list[dict[str, str]] = []
    errors: list[str] = []
    try:
        for task in asyncio.as_completed(tasks):
            try:
                results.append(await task)
            except HTTPException as exc:
                errors.append(str(exc.detail))
            except Exception as exc:
                errors.append(str(exc))
    finally:
        for task in tasks:
            if not task.done():
                task.cancel()
        await asyncio.gather(*tasks, return_exceptions=True)

    if not results:
        raise HTTPException(status_code=502, detail=f"all agent calls failed: {' | '.join(errors)[:300]}")
    return results


def _coerce_llm_chunk_text(content: Any) -> str:
    if isinstance(content, str):
        return content
    if isinstance(content, list):
        chunks: list[str] = []
        for item in content:
            if isinstance(item, str):
                chunks.append(item)
                continue
            if isinstance(item, dict):
                text = item.get("text")
                if isinstance(text, str):
                    chunks.append(text)
        return "".join(chunks)
    return ""


def _to_langchain_messages(context_messages: list[dict[str, str]]) -> list[Any]:
    try:
        from langchain_core.messages import AIMessage, HumanMessage, SystemMessage
    except Exception as exc:  # pragma: no cover - runtime dependency check
        raise HTTPException(status_code=500, detail=f"langchain dependency unavailable: {exc}") from exc

    mapped_messages: list[Any] = []
    for message in context_messages:
        role = str(message.get("role") or "user")
        content = str(message.get("content") or "").strip()
        if not content:
            continue
        if role == "system":
            mapped_messages.append(SystemMessage(content=content))
        elif role == "assistant":
            mapped_messages.append(AIMessage(content=content))
        else:
            mapped_messages.append(HumanMessage(content=content))
    return mapped_messages


async def _stream_chat_with_langchain(context_messages: list[dict[str, str]]) -> AsyncIterator[str]:
    try:
        from langchain_ollama import ChatOllama
    except Exception as exc:  # pragma: no cover - runtime dependency check
        raise HTTPException(status_code=500, detail=f"langchain ollama dependency unavailable: {exc}") from exc

    messages = _to_langchain_messages(context_messages)
    if not messages:
        raise HTTPException(status_code=400, detail="at least one non-empty message is required")

    llm = ChatOllama(
        model=OLLAMA_MODEL,
        base_url=OLLAMA_BASE_URL,
        keep_alive=_normalized_keep_alive(),
        temperature=0,
    )

    async for chunk in llm.astream(messages):
        if isinstance(chunk, str):
            text = chunk
        else:
            text = _coerce_llm_chunk_text(getattr(chunk, "content", ""))
        if text:
            yield text


def _superdoc_mcp_rpc(method: str, params: dict[str, Any], *, include_session: bool = True) -> dict[str, Any]:
    global superdoc_mcp_session_id

    timeout = httpx.Timeout(timeout=30.0, connect=5.0)
    request_body = {
        "jsonrpc": "2.0",
        "id": int(time.time() * 1000) % 1_000_000_000,
        "method": method,
        "params": params,
    }
    headers = {
        "accept": "application/json, text/event-stream",
        "content-type": "application/json",
    }
    if include_session and superdoc_mcp_session_id:
        headers["mcp-session-id"] = superdoc_mcp_session_id

    try:
        with httpx.Client(timeout=timeout) as http_client:
            response = http_client.post(SUPERDOC_MCP_URL, json=request_body, headers=headers)
    except httpx.HTTPError as exc:
        raise HTTPException(status_code=502, detail=f"superdoc mcp unavailable: {exc}") from exc

    if response.status_code >= 400:
        if include_session and response.status_code == 400 and superdoc_mcp_session_id:
            # Session may have expired in the MCP adapter; retry once with a fresh init path.
            superdoc_mcp_session_id = None
            return _superdoc_mcp_rpc(method, params, include_session=False)
        raise HTTPException(status_code=502, detail=f"superdoc mcp http error: {response.status_code} {response.text[:240]}")

    returned_session = response.headers.get("mcp-session-id")
    if returned_session:
        superdoc_mcp_session_id = returned_session

    parsed = _parse_streamable_http_body(response.text)
    if parsed.get("error"):
        message = parsed["error"].get("message") if isinstance(parsed["error"], dict) else str(parsed["error"])
        if include_session and superdoc_mcp_session_id and "session" in str(message).lower():
            superdoc_mcp_session_id = None
            return _superdoc_mcp_rpc(method, params, include_session=False)
        raise HTTPException(status_code=502, detail=f"superdoc mcp error: {message}")
    return parsed


def _ensure_superdoc_mcp_initialized() -> None:
    global superdoc_mcp_session_id
    if superdoc_mcp_session_id:
        return

    _superdoc_mcp_rpc(
        "initialize",
        {
            "protocolVersion": "2024-11-05",
            "capabilities": {},
            "clientInfo": {"name": "scholarflow-backend", "version": "0.1.0"},
        },
        include_session=False,
    )
    try:
        _superdoc_mcp_rpc("notifications/initialized", {}, include_session=True)
    except HTTPException:
        # Some gateways reject explicit initialized notifications; session init still stands.
        pass


def _superdoc_call_tool(action: str, arguments: dict[str, Any]) -> dict[str, Any]:
    _ensure_superdoc_mcp_initialized()
    tool_name = _resolve_superdoc_tool_name(action)
    payload = _superdoc_mcp_rpc(
        "tools/call",
        {"name": tool_name, "arguments": arguments},
        include_session=True,
    )
    result = payload.get("result")
    if not isinstance(result, dict):
        raise HTTPException(status_code=502, detail="superdoc mcp invalid tool result envelope")

    if result.get("isError"):
        tool_text = ""
        for item in result.get("content", []):
            if isinstance(item, dict):
                tool_text += str(item.get("text") or "")
        raise HTTPException(status_code=502, detail=f"superdoc tool error: {tool_text.strip()[:220]}")

    text_payload = ""
    for item in result.get("content", []):
        if isinstance(item, dict):
            text_payload += str(item.get("text") or "")

    text_payload = text_payload.strip()
    parsed: Any = None
    if text_payload:
        try:
            parsed = json.loads(text_payload)
        except json.JSONDecodeError:
            parsed = text_payload

    return {"raw": result, "text": text_payload, "parsed": parsed}


def _extract_json_object(text: str) -> dict[str, Any] | None:
    start = text.find("{")
    while start != -1:
        depth = 0
        for idx in range(start, len(text)):
            char = text[idx]
            if char == "{":
                depth += 1
            elif char == "}":
                depth -= 1
                if depth == 0:
                    candidate = text[start : idx + 1]
                    try:
                        parsed = json.loads(candidate)
                    except json.JSONDecodeError:
                        break
                    if isinstance(parsed, dict):
                        return parsed
                    break
        start = text.find("{", start + 1)
    return None


def _default_docx_draft(seed_text: str) -> tuple[str, dict[str, Any], str]:
    clean_seed = " ".join(seed_text.strip().split())
    title = "Untitled Draft"
    if clean_seed:
        title = _truncate_title(clean_seed)
    json_data = {
        "title": title,
        "author": {"name": "ScholarFlow AI"},
        "abstract": clean_seed or "Initial auto-generated draft from the latest conversation.",
        "body": clean_seed or "Add findings, evidence, and next steps from the conversation.",
    }
    template_xml = (
        "<doc>\n"
        "  <h1>{{title}}</h1>\n"
        "  <p><strong>Author:</strong> {{author.name}}</p>\n"
        "  <h2>Abstract</h2>\n"
        "  <p>{{abstract}}</p>\n"
        "  <h2>Draft</h2>\n"
        "  <p>{{body}}</p>\n"
        "</doc>"
    )
    filename = f"{re.sub(r'[^a-z0-9]+', '-', title.lower()).strip('-') or 'scholarflow-draft'}.docx"
    return template_xml, json_data, filename


def _normalize_docx_plan(payload: dict[str, Any], seed_text: str) -> tuple[str, dict[str, Any], str]:
    template_xml, json_data, filename = _default_docx_draft(seed_text)

    title = str(payload.get("title") or json_data["title"]).strip() or json_data["title"]
    author_name = "ScholarFlow AI"
    author = payload.get("author")
    if isinstance(author, dict):
        author_name = str(author.get("name") or author_name).strip() or author_name
    elif isinstance(author, str):
        author_name = author.strip() or author_name

    abstract = str(payload.get("abstract") or json_data["abstract"]).strip() or json_data["abstract"]
    body = str(payload.get("body") or "").strip()

    sections = payload.get("sections")
    if isinstance(sections, list):
        rendered_sections: list[str] = []
        for section in sections:
            if not isinstance(section, dict):
                continue
            heading = str(section.get("heading") or "").strip()
            content = str(section.get("content") or "").strip()
            if heading and content:
                rendered_sections.append(f"{heading}: {content}")
            elif content:
                rendered_sections.append(content)
        if rendered_sections:
            body = "\n".join(rendered_sections)

    if not body:
        body = json_data["body"]

    json_data = {
        "title": title,
        "author": {"name": author_name},
        "abstract": abstract,
        "body": body,
    }
    suggested_filename = str(payload.get("filename") or "").strip()
    if suggested_filename:
        filename = suggested_filename
    if not filename.lower().endswith(".docx"):
        filename = f"{filename}.docx"
    return template_xml, json_data, filename


async def _generate_docx_draft_with_ollama(
    seed_text: str,
    source_messages: list[dict[str, str]],
    current_template_xml: str | None = None,
    current_json_data: dict[str, Any] | None = None,
) -> tuple[str, dict[str, Any], str]:
    system_prompt = (
        "You are the ScholarFlow DOCX drafting tool. "
        "Return only JSON with keys: title, author, abstract, sections (array of {heading,content}), filename."
    )
    transcript_lines = [f"{item['role']}: {item['content']}" for item in source_messages if item.get("content")]
    transcript = "\n".join(transcript_lines[-16:])

    user_prompt = (
        "Create a professional DOCX draft plan from this conversation transcript. "
        "Keep abstract concise and sections practical. "
        "When current draft data is provided, treat this as an edit, not a reset.\n\n"
        f"Seed focus: {seed_text or 'general summary'}\n\n"
        f"Current template xml:\n{current_template_xml or '(none)'}\n\n"
        f"Current json data:\n{json.dumps(current_json_data or {}, ensure_ascii=True)}\n\n"
        f"Transcript:\n{transcript}"
    )

    payload = {
        "model": OLLAMA_MODEL,
        "stream": False,
        "keep_alive": _normalized_keep_alive(),
        "messages": [
            {"role": "system", "content": system_prompt},
            {"role": "user", "content": user_prompt},
        ],
    }

    timeout = httpx.Timeout(timeout=60.0, connect=10.0)
    async with httpx.AsyncClient(timeout=timeout) as http_client:
        response = await http_client.post(f"{OLLAMA_BASE_URL}/api/chat", json=payload)
        response.raise_for_status()
        response_payload = response.json()

    message_content = str(response_payload.get("message", {}).get("content", "")).strip()
    parsed = _extract_json_object(message_content)
    if not parsed:
        if current_template_xml and current_json_data:
            filename_seed = _truncate_title(seed_text or str(current_json_data.get("title") or "document"))
            safe_filename = f"{re.sub(r'[^a-z0-9]+', '-', filename_seed.lower()).strip('-') or 'document'}.docx"
            return current_template_xml, current_json_data, safe_filename
        return _default_docx_draft(seed_text)
    return _normalize_docx_plan(parsed, seed_text)


async def _generate_docx_draft_with_crewai(
    seed_text: str,
    source_messages: list[dict[str, str]],
    current_template_xml: str | None = None,
    current_json_data: dict[str, Any] | None = None,
) -> tuple[str, dict[str, Any], str]:
    return await asyncio.to_thread(
        _run_docx_generation_crew,
        seed_text,
        source_messages,
        current_template_xml,
        current_json_data,
    )


async def _generate_chat_with_crewai(context_messages: list[dict[str, str]]) -> list[dict[str, str]]:
    # Backwards-compatible function name; behavior is now parallel independent ChatDev-style agents.
    return await _generate_chatdev_like_responses(context_messages)


def _hash_password(password: str, salt: str) -> str:
    digest = hashlib.pbkdf2_hmac("sha256", password.encode("utf-8"), salt.encode("utf-8"), 120_000)
    return digest.hex()


def _verify_password(password: str, salt: str, expected_hash: str) -> bool:
    return hmac.compare_digest(_hash_password(password, salt), expected_hash)


def _scopes_for_role(role: str) -> list[str]:
    scopes = ["proposal:read", "proposal:write"]
    if role == "enterprise":
        scopes.append("plagiarism:access")
    return scopes


def _issue_tokens_for_user(user: dict[str, Any]) -> dict[str, Any]:
    now = datetime.now(timezone.utc)
    access_payload = {
        "sub": user["id"],
        "email": user["email"],
        "role": user["role"],
        "scopes": _scopes_for_role(user["role"]),
        "iat": int(now.timestamp()),
        "exp": int((now + timedelta(seconds=ACCESS_TTL_SECONDS)).timestamp()),
    }
    access_token = jwt.encode(access_payload, JWT_SECRET, algorithm=JWT_ALGORITHM)

    refresh_token = str(uuid.uuid4())
    refresh_tokens_by_user[user["id"]] = refresh_token
    return {
        "user": user,
        "access_token": access_token,
        "expires_in": ACCESS_TTL_SECONDS,
        "refresh_token": refresh_token,
        "token_type": "bearer",
    }


def _merge_crdt_operation(doc_id: str, operation: CrdtOperation) -> dict[str, Any]:
    state = doc_crdt_state[doc_id]
    if operation.op_id in state["applied_operation_ids"]:
        return {"doc_id": doc_id, "clock": state["clock"], "text": state["text"], "cursor": state["cursor"]}

    state["operations"].append(operation.model_dump())
    state["applied_operation_ids"].add(operation.op_id)
    ordered_ops = sorted(state["operations"], key=lambda op: (op["timestamp"], op["actor_id"], op["op_id"]))

    merged_text = ""
    cursor_map: dict[str, int] = {}
    max_ts = state["clock"]
    for op in ordered_ops:
        if op["text_patch"]:
            merged_text = f'{merged_text}{op["text_patch"]}'
        cursor_map[op["actor_id"]] = op["cursor_position"]
        max_ts = max(max_ts, op["timestamp"])

    state["clock"] = max_ts + 1
    state["text"] = merged_text
    state["cursor"] = cursor_map
    return {"doc_id": doc_id, "clock": state["clock"], "text": state["text"], "cursor": state["cursor"]}
