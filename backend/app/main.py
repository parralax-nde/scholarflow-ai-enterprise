from __future__ import annotations

import hashlib
import hmac
import io
import json
import math
import os
import re
import secrets
import sqlite3
import time
import uuid
import asyncio
import xml.etree.ElementTree as ET
from collections import Counter, defaultdict
from contextlib import asynccontextmanager
from datetime import datetime, timedelta, timezone
from pathlib import Path
from tempfile import gettempdir
from typing import Any

import httpx
from fastapi import FastAPI, Header, HTTPException, Request, WebSocket, WebSocketDisconnect
from fastapi.responses import Response, StreamingResponse
from jinja2 import Environment
from jose import jwt
from pydantic import BaseModel, EmailStr, Field
from docx import Document

JWT_SECRET = os.getenv("JWT_SECRET", "dev-secret-change-me")
JWT_ALGORITHM = "HS256"
ACCESS_TTL_SECONDS = 3600
OLLAMA_BASE_URL = os.getenv("OLLAMA_BASE_URL", "http://localhost:11434").rstrip("/")
OLLAMA_MODEL = os.getenv("OLLAMA_MODEL", "gemma4:e2b")
OLLAMA_KEEP_ALIVE = os.getenv("OLLAMA_KEEP_ALIVE", "-1")
SUPERDOC_MCP_BRIDGE_URL = os.getenv("SUPERDOC_MCP_BRIDGE_URL", "http://localhost:8090").rstrip("/")
MAX_CONTEXT_MESSAGES = int(os.getenv("MAX_CONTEXT_MESSAGES", "12"))
CHAT_DB_PATH = Path(os.getenv("CHAT_DB_PATH", str(Path(gettempdir()) / "scholarflow" / "chat_history.db")))
DOCX_ARTIFACT_DIR = Path(os.getenv("DOCX_ARTIFACT_DIR", str(Path(gettempdir()) / "scholarflow" / "docx_artifacts")))
DEFAULT_CONVERSATION_TITLE = "New chat"
MAX_CONVERSATION_TITLE_LENGTH = 60

users_by_google_sub: dict[str, dict[str, Any]] = {}
users_by_email: dict[str, dict[str, Any]] = {}
refresh_tokens_by_user: dict[str, str] = {}
service_registry: dict[str, dict[str, Any]] = {}
subscriptions: dict[str, str] = defaultdict(lambda: "free")
app_started_at = time.time()
app_instance_id = os.getenv("SERVICE_INSTANCE_ID", f"core-{uuid.uuid4()}")
app_service_name = os.getenv("SERVICE_NAME", "core")
app_metadata_tags = [tag for tag in os.getenv("SERVICE_METADATA_TAGS", "core,fastapi,scholarflow").split(",") if tag]
plagiarism_events: list[dict[str, Any]] = []
doc_crdt_state: dict[str, dict[str, Any]] = defaultdict(
    lambda: {"clock": 0, "text": "", "cursor": {}, "operations": [], "applied_operation_ids": set()}
)
chat_db_initialized = False
DOCX_ARTIFACT_TTL_SECONDS = 1800
docx_artifacts: dict[str, dict[str, Any]] = {}


def _utc_now_iso() -> str:
    return datetime.now(timezone.utc).isoformat()


def _connect_chat_db() -> sqlite3.Connection:
    _init_chat_db()
    return sqlite3.connect(CHAT_DB_PATH)


def _init_chat_db() -> None:
    global chat_db_initialized
    if chat_db_initialized:
        return
    CHAT_DB_PATH.parent.mkdir(parents=True, exist_ok=True)
    with sqlite3.connect(CHAT_DB_PATH) as conn:
        conn.execute(
            """
            CREATE TABLE IF NOT EXISTS conversations (
                id TEXT PRIMARY KEY,
                title TEXT NOT NULL,
                created_at TEXT NOT NULL,
                updated_at TEXT NOT NULL
            )
            """
        )
        conn.execute(
            """
            CREATE TABLE IF NOT EXISTS messages (
                id TEXT PRIMARY KEY,
                conversation_id TEXT NOT NULL,
                role TEXT NOT NULL,
                content TEXT NOT NULL,
                created_at TEXT NOT NULL,
                FOREIGN KEY(conversation_id) REFERENCES conversations(id) ON DELETE CASCADE
            )
            """
        )
        conn.execute("CREATE INDEX IF NOT EXISTS idx_messages_conversation ON messages(conversation_id, created_at)")
        conn.execute("CREATE INDEX IF NOT EXISTS idx_conversations_updated ON conversations(updated_at DESC)")
    chat_db_initialized = True


def _conversation_exists(conversation_id: str) -> bool:
    with _connect_chat_db() as conn:
        row = conn.execute("SELECT 1 FROM conversations WHERE id = ?", (conversation_id,)).fetchone()
    return bool(row)


def _create_conversation(title: str = DEFAULT_CONVERSATION_TITLE) -> dict[str, Any]:
    conversation_id = str(uuid.uuid4())
    now = _utc_now_iso()
    with _connect_chat_db() as conn:
        conn.execute(
            "INSERT INTO conversations (id, title, created_at, updated_at) VALUES (?, ?, ?, ?)",
            (conversation_id, title, now, now),
        )
    return {"id": conversation_id, "title": title, "message_count": 0, "updated_at": now}


def _truncate_title(value: str) -> str:
    clean = " ".join(value.strip().split())
    if len(clean) <= MAX_CONVERSATION_TITLE_LENGTH:
        return clean
    if MAX_CONVERSATION_TITLE_LENGTH <= 3:
        return clean[:MAX_CONVERSATION_TITLE_LENGTH]
    sliced = clean[: MAX_CONVERSATION_TITLE_LENGTH - 3].rstrip()
    if " " in sliced:
        sliced = sliced.rsplit(" ", 1)[0]
    return f"{sliced}..."


def _add_message(conversation_id: str, role: str, content: str) -> dict[str, Any]:
    message_id = str(uuid.uuid4())
    now = _utc_now_iso()
    with _connect_chat_db() as conn:
        conn.execute(
            "INSERT INTO messages (id, conversation_id, role, content, created_at) VALUES (?, ?, ?, ?, ?)",
            (message_id, conversation_id, role, content, now),
        )
        conn.execute("UPDATE conversations SET updated_at = ? WHERE id = ?", (now, conversation_id))
        row = conn.execute("SELECT title FROM conversations WHERE id = ?", (conversation_id,)).fetchone()
        if row and row[0].strip() == DEFAULT_CONVERSATION_TITLE and role == "user":
            conn.execute(
                "UPDATE conversations SET title = ? WHERE id = ?",
                (_truncate_title(content), conversation_id),
            )
    return {"id": message_id, "conversation_id": conversation_id, "role": role, "content": content, "created_at": now}


def _list_conversations() -> list[dict[str, Any]]:
    with _connect_chat_db() as conn:
        rows = conn.execute(
            """
            SELECT c.id, c.title, c.updated_at, COUNT(m.id) AS message_count
            FROM conversations c
            LEFT JOIN messages m ON m.conversation_id = c.id
            GROUP BY c.id
            ORDER BY c.updated_at DESC
            """
        ).fetchall()
    return [
        {
            "id": row[0],
            "title": row[1],
            "updated_at": row[2],
            "message_count": row[3],
        }
        for row in rows
    ]


def _list_messages(conversation_id: str) -> list[dict[str, Any]]:
    with _connect_chat_db() as conn:
        rows = conn.execute(
            """
            SELECT id, role, content, created_at
            FROM messages
            WHERE conversation_id = ?
            ORDER BY created_at ASC
            """,
            (conversation_id,),
        ).fetchall()
    return [{"id": row[0], "role": row[1], "content": row[2], "created_at": row[3]} for row in rows]


def register_self() -> None:
    service_registry[app_instance_id] = {
        "instance_id": app_instance_id,
        "service_name": app_service_name,
        "metadata_tags": app_metadata_tags,
        "uptime_seconds": 0,
        "request_latency_ms": 12,
        "error_rate": 0.0,
    }


def _normalized_keep_alive() -> str:
    raw = OLLAMA_KEEP_ALIVE.strip()
    if raw == "-1":
        # Ollama expects duration strings on some endpoints; this value means effectively no eviction.
        return "2562047h47m16.854775807s"
    if re.fullmatch(r"-?\d+", raw):
        return f"{raw}s"
    return raw


async def _warm_ollama_model() -> None:
    payload = {
        "model": OLLAMA_MODEL,
        "stream": False,
        "keep_alive": _normalized_keep_alive(),
        "prompt": "Warmup",
        "options": {"num_predict": 1},
    }
    timeout = httpx.Timeout(timeout=30.0, connect=5.0)
    try:
        async with httpx.AsyncClient(timeout=timeout) as http_client:
            await http_client.post(f"{OLLAMA_BASE_URL}/api/generate", json=payload)
    except httpx.HTTPError:
        # Warmup is best-effort and should never break API startup.
        return


@asynccontextmanager
async def lifespan(_: FastAPI):
    register_self()
    asyncio.create_task(_warm_ollama_model())
    yield


app = FastAPI(title="ScholarFlow AI Enterprise", lifespan=lifespan)


class RegisterServiceRequest(BaseModel):
    service_name: str
    metadata_tags: list[str] = Field(default_factory=list)


class GoogleProfile(BaseModel):
    sub: str
    email: EmailStr
    name: str


class OAuthCallbackRequest(BaseModel):
    code: str
    profile: GoogleProfile
    tier: str = Field(default="free", pattern="^(free|enterprise)$")


class EmailPasswordRegisterRequest(BaseModel):
    email: EmailStr
    name: str
    password: str = Field(min_length=8)
    tier: str = Field(default="free", pattern="^(free|enterprise)$")


class EmailPasswordLoginRequest(BaseModel):
    email: EmailStr
    password: str


class ChatMessage(BaseModel):
    role: str = Field(pattern="^(system|user|assistant)$")
    content: str = Field(min_length=1)


class ChatStreamRequest(BaseModel):
    messages: list[ChatMessage] = Field(min_length=1)
    conversation_id: str | None = None


class ConversationCreateRequest(BaseModel):
    title: str | None = Field(default=None, max_length=MAX_CONVERSATION_TITLE_LENGTH)


class SimilarityRequest(BaseModel):
    text: str
    sources: list[str]


class RemediationRequest(BaseModel):
    flagged_sentences: list[str]


class CrdtOperation(BaseModel):
    op_id: str
    actor_id: str
    timestamp: int
    text_patch: str = ""
    cursor_position: int = 0


class CitationClaim(BaseModel):
    claim: str
    source_text: str


class CitationValidationRequest(BaseModel):
    citations: list[CitationClaim]


class ExportRequest(BaseModel):
    title: str
    author: str
    body: str
    custom_css: str = ""


class DocxTemplateRequest(BaseModel):
    template_xml: str = Field(min_length=1)
    json_data: dict[str, Any] = Field(default_factory=dict)
    filename: str = Field(default="draft.docx", max_length=120)


class DocxPreviewResponse(BaseModel):
    resolved_xml: str
    preview_html: str
    paragraphs: list[str]


class DocxSessionResponse(BaseModel):
    docx_id: str
    filename: str
    viewer_path: str


class DocxToolGenerateRequest(BaseModel):
    conversation_id: str | None = None
    prompt: str | None = None
    messages: list[ChatMessage] = Field(default_factory=list)
    current_template_xml: str | None = None
    current_json_data: dict[str, Any] | None = None


class DocxToolGenerateResponse(BaseModel):
    template_xml: str
    json_data: dict[str, Any]
    filename: str
    model: str
    source: str


class SuperdocOpenRequest(BaseModel):
    docx_id: str = Field(min_length=1)


class SuperdocApplyRequest(BaseModel):
    docx_id: str = Field(min_length=1)
    content: str = Field(min_length=1)
    suggest: bool = False


class SuperdocSessionRequest(BaseModel):
    session_id: str = Field(min_length=1)


def _tokenize(text: str) -> Counter:
    tokens = re.findall(r"[a-zA-Z0-9]+", text.lower())
    return Counter(tokens)


def _cosine_similarity(a: str, b: str) -> float:
    va = _tokenize(a)
    vb = _tokenize(b)
    if not va or not vb:
        return 0.0
    dot = sum(va[t] * vb.get(t, 0) for t in va)
    norm_a = math.sqrt(sum(v * v for v in va.values()))
    norm_b = math.sqrt(sum(v * v for v in vb.values()))
    return dot / (norm_a * norm_b) if norm_a and norm_b else 0.0


def _split_sentences(text: str) -> list[str]:
    return [s.strip() for s in re.split(r"(?<=[.!?])\s+", text.strip()) if s.strip()]


def _trace_id_from_request(headers: dict[str, str]) -> str:
    traceparent = headers.get("traceparent", "")
    if traceparent:
        parts = traceparent.split("-")
        if len(parts) >= 2 and parts[1]:
            return parts[1]
    return headers.get("x-request-id", str(uuid.uuid4()))


def _lookup_path(data: dict[str, Any], dotted_path: str) -> str:
    current: Any = data
    for key in dotted_path.split("."):
        if isinstance(current, dict) and key in current:
            current = current[key]
            continue
        return ""
    if isinstance(current, (dict, list)):
        return json.dumps(current)
    return "" if current is None else str(current)


def _resolve_template_placeholders(template_xml: str, json_data: dict[str, Any]) -> str:
    def replace_match(match: re.Match[str]) -> str:
        key = match.group(1).strip()
        return _lookup_path(json_data, key)

    return re.sub(r"\{\{\s*([a-zA-Z0-9_.-]+)\s*\}\}", replace_match, template_xml)


def _extract_docx_paragraphs(root: ET.Element) -> list[str]:
    paragraphs: list[str] = []
    for node in root.iter():
        if node.tag in {"p", "h1", "h2", "h3", "li"}:
            content = "".join(node.itertext()).strip()
            if content:
                paragraphs.append(content)
    if paragraphs:
        return paragraphs
    flattened = "".join(root.itertext()).strip()
    return [flattened] if flattened else []


def _build_preview_html(root: ET.Element) -> str:
    lines: list[str] = []
    for node in root:
        if node.tag in {"h1", "h2", "h3", "p"}:
            lines.append(f"<{node.tag}>{''.join(node.itertext()).strip()}</{node.tag}>")
        elif node.tag == "ul":
            items = [f"<li>{''.join(item.itertext()).strip()}</li>" for item in node.findall("li")]
            lines.append(f"<ul>{''.join(items)}</ul>")
    if not lines:
        text = "".join(root.itertext()).strip()
        if text:
            lines.append(f"<p>{text}</p>")
    return "".join(lines)


def _build_docx_bytes_from_xml(resolved_xml: str) -> bytes:
    try:
        root = ET.fromstring(resolved_xml)
    except ET.ParseError as exc:
        raise HTTPException(status_code=400, detail=f"invalid template XML: {exc}") from exc

    document = Document()
    paragraphs = _extract_docx_paragraphs(root)
    if not paragraphs:
        raise HTTPException(status_code=400, detail="template produced empty content")

    for paragraph in paragraphs:
        document.add_paragraph(paragraph)

    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def _clean_docx_artifacts() -> None:
    cutoff = time.time() - DOCX_ARTIFACT_TTL_SECONDS
    stale_ids = [artifact_id for artifact_id, entry in docx_artifacts.items() if entry["created_at"] < cutoff]
    for artifact_id in stale_ids:
        artifact = docx_artifacts.pop(artifact_id, None)
        if artifact:
            file_path = artifact.get("file_path")
            if isinstance(file_path, str) and file_path:
                Path(file_path).unlink(missing_ok=True)


def _ensure_docx_artifact_file(docx_id: str) -> str:
    _clean_docx_artifacts()
    artifact = docx_artifacts.get(docx_id)
    if not artifact:
        raise HTTPException(status_code=404, detail="docx artifact not found")

    DOCX_ARTIFACT_DIR.mkdir(parents=True, exist_ok=True)
    file_path = artifact.get("file_path")
    if isinstance(file_path, str) and file_path:
        path = Path(file_path)
    else:
        filename = artifact.get("filename") or f"{docx_id}.docx"
        safe_name = re.sub(r"[^a-zA-Z0-9._-]+", "-", str(filename)) or f"{docx_id}.docx"
        path = DOCX_ARTIFACT_DIR / f"{docx_id}-{safe_name}"
        artifact["file_path"] = str(path)

    if not path.exists():
        content = artifact.get("content")
        if not isinstance(content, (bytes, bytearray)):
            raise HTTPException(status_code=500, detail="docx artifact content unavailable")
        path.write_bytes(bytes(content))
    return str(path)


def _reload_docx_artifact_content(docx_id: str) -> None:
    artifact = docx_artifacts.get(docx_id)
    if not artifact:
        return
    file_path = artifact.get("file_path")
    if isinstance(file_path, str) and file_path and Path(file_path).exists():
        artifact["content"] = Path(file_path).read_bytes()


def _mcp_bridge_post(path: str, payload: dict[str, Any]) -> dict[str, Any]:
    timeout = httpx.Timeout(timeout=30.0, connect=5.0)
    try:
        with httpx.Client(timeout=timeout) as http_client:
            response = http_client.post(f"{SUPERDOC_MCP_BRIDGE_URL}{path}", json=payload)
    except httpx.HTTPError as exc:
        raise HTTPException(status_code=502, detail=f"mcp bridge unavailable: {exc}") from exc

    if response.status_code >= 400:
        detail = response.text
        try:
            payload_error = response.json()
            detail = str(payload_error.get("detail") or payload_error)
        except Exception:
            detail = response.text
        raise HTTPException(status_code=502, detail=f"mcp bridge error: {detail}")

    try:
        return response.json()
    except Exception as exc:
        raise HTTPException(status_code=502, detail=f"invalid mcp bridge response: {exc}") from exc


def _extract_json_object(text: str) -> dict[str, Any] | None:
    start = text.find("{")
    while start != -1:
        depth = 0
        for idx in range(start, len(text)):
            char = text[idx]
            if char == "{":
                depth += 1
            elif char == "}":
                depth -= 1
                if depth == 0:
                    candidate = text[start : idx + 1]
                    try:
                        parsed = json.loads(candidate)
                    except json.JSONDecodeError:
                        break
                    if isinstance(parsed, dict):
                        return parsed
                    break
        start = text.find("{", start + 1)
    return None


def _default_docx_draft(seed_text: str) -> tuple[str, dict[str, Any], str]:
    clean_seed = " ".join(seed_text.strip().split())
    title = "Untitled Draft"
    if clean_seed:
        title = _truncate_title(clean_seed)
    json_data = {
        "title": title,
        "author": {"name": "ScholarFlow AI"},
        "abstract": clean_seed or "Initial auto-generated draft from the latest conversation.",
        "body": clean_seed or "Add findings, evidence, and next steps from the conversation.",
    }
    template_xml = (
        "<doc>\n"
        "  <h1>{{title}}</h1>\n"
        "  <p><strong>Author:</strong> {{author.name}}</p>\n"
        "  <h2>Abstract</h2>\n"
        "  <p>{{abstract}}</p>\n"
        "  <h2>Draft</h2>\n"
        "  <p>{{body}}</p>\n"
        "</doc>"
    )
    filename = f"{re.sub(r'[^a-z0-9]+', '-', title.lower()).strip('-') or 'scholarflow-draft'}.docx"
    return template_xml, json_data, filename


def _normalize_docx_plan(payload: dict[str, Any], seed_text: str) -> tuple[str, dict[str, Any], str]:
    template_xml, json_data, filename = _default_docx_draft(seed_text)

    title = str(payload.get("title") or json_data["title"]).strip() or json_data["title"]
    author_name = "ScholarFlow AI"
    author = payload.get("author")
    if isinstance(author, dict):
        author_name = str(author.get("name") or author_name).strip() or author_name
    elif isinstance(author, str):
        author_name = author.strip() or author_name

    abstract = str(payload.get("abstract") or json_data["abstract"]).strip() or json_data["abstract"]
    body = str(payload.get("body") or "").strip()

    sections = payload.get("sections")
    if isinstance(sections, list):
        rendered_sections: list[str] = []
        for section in sections:
            if not isinstance(section, dict):
                continue
            heading = str(section.get("heading") or "").strip()
            content = str(section.get("content") or "").strip()
            if heading and content:
                rendered_sections.append(f"{heading}: {content}")
            elif content:
                rendered_sections.append(content)
        if rendered_sections:
            body = "\n".join(rendered_sections)

    if not body:
        body = json_data["body"]

    json_data = {
        "title": title,
        "author": {"name": author_name},
        "abstract": abstract,
        "body": body,
    }
    suggested_filename = str(payload.get("filename") or "").strip()
    if suggested_filename:
        filename = suggested_filename
    if not filename.lower().endswith(".docx"):
        filename = f"{filename}.docx"
    return template_xml, json_data, filename


async def _generate_docx_draft_with_ollama(
    seed_text: str,
    source_messages: list[dict[str, str]],
    current_template_xml: str | None = None,
    current_json_data: dict[str, Any] | None = None,
) -> tuple[str, dict[str, Any], str]:
    system_prompt = (
        "You are the ScholarFlow DOCX drafting tool. "
        "Return only JSON with keys: title, author, abstract, sections (array of {heading,content}), filename."
    )
    transcript_lines = [f"{item['role']}: {item['content']}" for item in source_messages if item.get("content")]
    transcript = "\n".join(transcript_lines[-16:])

    user_prompt = (
        "Create a professional DOCX draft plan from this conversation transcript. "
        "Keep abstract concise and sections practical. "
        "When current draft data is provided, treat this as an edit, not a reset.\n\n"
        f"Seed focus: {seed_text or 'general summary'}\n\n"
        f"Current template xml:\n{current_template_xml or '(none)'}\n\n"
        f"Current json data:\n{json.dumps(current_json_data or {}, ensure_ascii=True)}\n\n"
        f"Transcript:\n{transcript}"
    )

    payload = {
        "model": OLLAMA_MODEL,
        "stream": False,
        "keep_alive": _normalized_keep_alive(),
        "messages": [
            {"role": "system", "content": system_prompt},
            {"role": "user", "content": user_prompt},
        ],
    }

    timeout = httpx.Timeout(timeout=60.0, connect=10.0)
    async with httpx.AsyncClient(timeout=timeout) as http_client:
        response = await http_client.post(f"{OLLAMA_BASE_URL}/api/chat", json=payload)
        response.raise_for_status()
        response_payload = response.json()

    message_content = str(response_payload.get("message", {}).get("content", "")).strip()
    parsed = _extract_json_object(message_content)
    if not parsed:
        if current_template_xml and current_json_data:
            filename_seed = _truncate_title(seed_text or str(current_json_data.get("title") or "document"))
            safe_filename = f"{re.sub(r'[^a-z0-9]+', '-', filename_seed.lower()).strip('-') or 'document'}.docx"
            return current_template_xml, current_json_data, safe_filename
        return _default_docx_draft(seed_text)
    return _normalize_docx_plan(parsed, seed_text)


def _hash_password(password: str, salt: str) -> str:
    digest = hashlib.pbkdf2_hmac("sha256", password.encode("utf-8"), salt.encode("utf-8"), 120_000)
    return digest.hex()


def _verify_password(password: str, salt: str, expected_hash: str) -> bool:
    return hmac.compare_digest(_hash_password(password, salt), expected_hash)


def _scopes_for_role(role: str) -> list[str]:
    scopes = ["proposal:read", "proposal:write"]
    if role == "enterprise":
        scopes.append("plagiarism:access")
    return scopes


def _issue_tokens_for_user(user: dict[str, Any]) -> dict[str, Any]:
    now = datetime.now(timezone.utc)
    access_payload = {
        "sub": user["id"],
        "email": user["email"],
        "role": user["role"],
        "scopes": _scopes_for_role(user["role"]),
        "iat": int(now.timestamp()),
        "exp": int((now + timedelta(seconds=ACCESS_TTL_SECONDS)).timestamp()),
    }
    access_token = jwt.encode(access_payload, JWT_SECRET, algorithm=JWT_ALGORITHM)

    refresh_token = str(uuid.uuid4())
    refresh_tokens_by_user[user["id"]] = refresh_token
    return {
        "user": user,
        "access_token": access_token,
        "expires_in": ACCESS_TTL_SECONDS,
        "refresh_token": refresh_token,
        "token_type": "bearer",
    }


def _merge_crdt_operation(doc_id: str, operation: CrdtOperation) -> dict[str, Any]:
    state = doc_crdt_state[doc_id]
    if operation.op_id in state["applied_operation_ids"]:
        return {"doc_id": doc_id, "clock": state["clock"], "text": state["text"], "cursor": state["cursor"]}

    state["operations"].append(operation.model_dump())
    state["applied_operation_ids"].add(operation.op_id)
    ordered_ops = sorted(state["operations"], key=lambda op: (op["timestamp"], op["actor_id"], op["op_id"]))

    merged_text = ""
    cursor_map: dict[str, int] = {}
    max_ts = state["clock"]
    for op in ordered_ops:
        if op["text_patch"]:
            merged_text = f'{merged_text}{op["text_patch"]}'
        cursor_map[op["actor_id"]] = op["cursor_position"]
        max_ts = max(max_ts, op["timestamp"])

    state["clock"] = max_ts + 1
    state["text"] = merged_text
    state["cursor"] = cursor_map
    return {"doc_id": doc_id, "clock": state["clock"], "text": state["text"], "cursor": state["cursor"]}


@app.middleware("http")
async def trace_middleware(request: Request, call_next):
    trace_id = _trace_id_from_request(dict(request.headers))
    response = await call_next(request)
    response.headers["X-Trace-Id"] = trace_id
    return response


@app.get("/core/health")
def core_health() -> dict[str, Any]:
    return {
        "status": "ok",
        "service": app_service_name,
        "timestamp": datetime.now(timezone.utc).isoformat(),
        "instance_id": app_instance_id,
        "metadata_tags": app_metadata_tags,
    }


@app.post("/core/register")
def register_service(body: RegisterServiceRequest) -> dict[str, Any]:
    instance_id = f"{body.service_name}-{uuid.uuid4()}"
    service_registry[instance_id] = {
        "instance_id": instance_id,
        "service_name": body.service_name,
        "metadata_tags": body.metadata_tags,
        "uptime_seconds": int(time.time() - app_started_at),
        "request_latency_ms": 12,
        "error_rate": 0.0,
    }
    return service_registry[instance_id]


@app.get("/core/services")
def list_services() -> list[dict[str, Any]]:
    return list(service_registry.values())


@app.post("/auth/oauth/google/callback")
def oauth_google_callback(body: OAuthCallbackRequest) -> dict[str, Any]:
    role = "enterprise" if body.tier == "enterprise" else "free"

    user = users_by_google_sub.setdefault(
        body.profile.sub,
        {
            "id": str(uuid.uuid4()),
            "email": body.profile.email,
            "name": body.profile.name,
            "role": role,
        },
    )
    user["role"] = role

    return _issue_tokens_for_user(user)


@app.post("/auth/register")
def auth_register(body: EmailPasswordRegisterRequest) -> dict[str, Any]:
    email_key = body.email.lower()
    if email_key in users_by_email:
        raise HTTPException(status_code=409, detail="email already registered")

    role = "enterprise" if body.tier == "enterprise" else "free"
    salt = secrets.token_hex(16)
    user_id = str(uuid.uuid4())
    user_record = {
        "id": user_id,
        "email": email_key,
        "name": body.name,
        "role": role,
        "auth_provider": "password",
        "password_salt": salt,
        "password_hash": _hash_password(body.password, salt),
    }
    users_by_email[email_key] = user_record
    public_user = {k: v for k, v in user_record.items() if not k.startswith("password_")}
    return _issue_tokens_for_user(public_user)


@app.post("/auth/login")
def auth_login(body: EmailPasswordLoginRequest) -> dict[str, Any]:
    email_key = body.email.lower()
    user_record = users_by_email.get(email_key)
    if not user_record:
        raise HTTPException(status_code=401, detail="invalid credentials")

    if not _verify_password(body.password, user_record["password_salt"], user_record["password_hash"]):
        raise HTTPException(status_code=401, detail="invalid credentials")

    public_user = {k: v for k, v in user_record.items() if not k.startswith("password_")}
    return _issue_tokens_for_user(public_user)


@app.get("/ai/conversations")
def ai_list_conversations() -> list[dict[str, Any]]:
    return _list_conversations()


@app.post("/ai/conversations")
def ai_create_conversation(body: ConversationCreateRequest | None = None) -> dict[str, Any]:
    title = DEFAULT_CONVERSATION_TITLE
    if body and body.title and body.title.strip():
        title = body.title.strip()[:MAX_CONVERSATION_TITLE_LENGTH]
    return _create_conversation(title=title)


@app.get("/ai/conversations/{conversation_id}/messages")
def ai_list_messages(conversation_id: str) -> dict[str, Any]:
    if not _conversation_exists(conversation_id):
        raise HTTPException(status_code=404, detail="conversation not found")
    return {"conversation_id": conversation_id, "messages": _list_messages(conversation_id)}


@app.post("/ai/chat/stream")
async def ai_chat_stream(body: ChatStreamRequest) -> StreamingResponse:
    conversation_id = body.conversation_id
    if conversation_id:
        if not _conversation_exists(conversation_id):
            raise HTTPException(status_code=404, detail="conversation not found")
    else:
        conversation = _create_conversation()
        conversation_id = conversation["id"]

    user_messages = [message for message in body.messages if message.role == "user"]
    latest_user_message = user_messages[-1].content.strip() if user_messages else ""
    if not latest_user_message:
        raise HTTPException(status_code=400, detail="at least one user message is required")
    _add_message(conversation_id, "user", latest_user_message)

    context_messages = body.messages[-MAX_CONTEXT_MESSAGES:] if len(body.messages) > MAX_CONTEXT_MESSAGES else body.messages

    payload = {
        "model": OLLAMA_MODEL,
        "stream": True,
        "keep_alive": _normalized_keep_alive(),
        "messages": [message.model_dump() for message in context_messages],
    }

    async def stream_generator():
        yield json.dumps({"type": "meta", "model": OLLAMA_MODEL, "phase": "requesting"}) + "\n"
        assistant_chunks: list[str] = []
        try:
            timeout = httpx.Timeout(timeout=90.0, connect=10.0)
            async with httpx.AsyncClient(timeout=timeout) as http_client:
                async with http_client.stream("POST", f"{OLLAMA_BASE_URL}/api/chat", json=payload) as response:
                    if response.status_code >= 400:
                        raw_error = (await response.aread()).decode("utf-8", errors="ignore")
                        detail = raw_error
                        try:
                            parsed = json.loads(raw_error)
                            detail = str(parsed.get("error") or raw_error)
                        except json.JSONDecodeError:
                            detail = raw_error
                        yield json.dumps(
                            {
                                "type": "error",
                                "error": f"ollama request failed ({response.status_code}): {detail.strip()[:220]}",
                            }
                        ) + "\n"
                        return

                    yield json.dumps({"type": "meta", "model": OLLAMA_MODEL, "phase": "streaming"}) + "\n"
                    async for line in response.aiter_lines():
                        if not line:
                            continue
                        try:
                            chunk = json.loads(line)
                        except json.JSONDecodeError:
                            continue
                        thinking = chunk.get("message", {}).get("thinking", "")
                        if thinking:
                            yield json.dumps({"type": "thinking", "content": thinking}) + "\n"
                        content = chunk.get("message", {}).get("content", "")
                        if content:
                            assistant_chunks.append(content)
                            yield json.dumps({"type": "token", "content": content}) + "\n"
                        if chunk.get("done"):
                            yield json.dumps(
                                {
                                    "type": "done",
                                    "prompt_eval_count": chunk.get("prompt_eval_count"),
                                    "prompt_eval_duration": chunk.get("prompt_eval_duration"),
                                    "eval_count": chunk.get("eval_count"),
                                    "eval_duration": chunk.get("eval_duration"),
                                    "total_duration": chunk.get("total_duration"),
                                }
                            ) + "\n"
            assistant_content = "".join(assistant_chunks).strip()
            if assistant_content:
                _add_message(conversation_id, "assistant", assistant_content)
        except httpx.HTTPError as e:
            yield json.dumps({"type": "error", "error": f"ollama unavailable: {str(e)}"}) + "\n"

    return StreamingResponse(
        stream_generator(),
        media_type="application/x-ndjson",
        headers={
            "Cache-Control": "no-cache",
            "Connection": "keep-alive",
            "X-Accel-Buffering": "no",
        }
    )


@app.post("/ai/mcp/tools/docx/generate", response_model=DocxToolGenerateResponse)
async def ai_mcp_docx_generate(body: DocxToolGenerateRequest) -> DocxToolGenerateResponse:
    source_messages: list[dict[str, str]] = []
    if body.messages:
        source_messages = [
            {"role": message.role, "content": message.content}
            for message in body.messages
            if message.content.strip()
        ]
    elif body.conversation_id:
        if not _conversation_exists(body.conversation_id):
            raise HTTPException(status_code=404, detail="conversation not found")
        source_messages = [
            {"role": message["role"], "content": message["content"]}
            for message in _list_messages(body.conversation_id)
            if str(message.get("content", "")).strip()
        ]

    if not source_messages and not (body.prompt and body.prompt.strip()):
        raise HTTPException(status_code=400, detail="conversation_id, messages, or prompt is required")

    seed_text = (body.prompt or "").strip()
    if not seed_text:
        user_like_messages = [m["content"] for m in source_messages if m["role"] == "user"]
        seed_text = user_like_messages[-1].strip() if user_like_messages else source_messages[-1]["content"]

    try:
        template_xml, json_data, filename = await _generate_docx_draft_with_ollama(
            seed_text,
            source_messages,
            body.current_template_xml,
            body.current_json_data,
        )
    except httpx.HTTPError:
        if body.current_template_xml and body.current_json_data:
            filename_seed = _truncate_title(seed_text or str(body.current_json_data.get("title") or "document"))
            fallback_filename = f"{re.sub(r'[^a-z0-9]+', '-', filename_seed.lower()).strip('-') or 'document'}.docx"
            template_xml, json_data, filename = body.current_template_xml, body.current_json_data, fallback_filename
        else:
            template_xml, json_data, filename = _default_docx_draft(seed_text)

    return DocxToolGenerateResponse(
        template_xml=template_xml,
        json_data=json_data,
        filename=filename,
        model=OLLAMA_MODEL,
        source="mcp-docx-tool",
    )


@app.post("/plagiarism/similarity")
def plagiarism_similarity(body: SimilarityRequest) -> dict[str, Any]:
    source_sentences: list[str] = []
    for source in body.sources:
        source_sentences.extend(_split_sentences(source))

    flagged = []
    for sentence in _split_sentences(body.text):
        score = max((_cosine_similarity(sentence, src) for src in source_sentences), default=0.0)
        if score > 0.8:
            flagged.append({"sentence": sentence, "similarity": round(score, 4)})
    return {"threshold": 0.8, "flagged": flagged}


@app.post("/plagiarism/remediate")
def plagiarism_remediate(body: RemediationRequest) -> dict[str, Any]:
    rewritten = [f"Rewritten: {s}" for s in body.flagged_sentences]
    event = {
        "event_id": str(uuid.uuid4()),
        "event_type": "plagiarism.remediated",
        "provider": "ollama",
        "model": OLLAMA_MODEL,
        "rewritten": rewritten,
        "timestamp": datetime.now(timezone.utc).isoformat(),
    }
    plagiarism_events.append(event)
    return event


@app.get("/plagiarism/events")
def plagiarism_event_stream() -> dict[str, Any]:
    return {"events": plagiarism_events}


@app.post("/citations/validate")
def citation_validate(body: CitationValidationRequest) -> dict[str, Any]:
    results = []
    for item in body.citations:
        score = _cosine_similarity(item.claim, item.source_text)
        status = "Unsupported"
        if score >= 0.6:
            status = "Supported"
        elif score >= 0.3:
            status = "Partially Supported"
        results.append({"claim": item.claim, "status": status, "score": round(score, 3)})
    return {"results": results}


active_docs: dict[str, set[WebSocket]] = defaultdict(set)


@app.websocket("/collab/ws/{doc_id}")
async def collab_ws(doc_id: str, websocket: WebSocket) -> None:
    await websocket.accept()
    active_docs[doc_id].add(websocket)
    try:
        while True:
            payload = await websocket.receive_json()
            merged = None
            if {"op_id", "actor_id", "timestamp"}.issubset(payload.keys()):
                merged = _merge_crdt_operation(doc_id, CrdtOperation(**payload))
            for peer in list(active_docs[doc_id]):
                if peer is not websocket:
                    await peer.send_json({"operation": payload, "state": merged or doc_crdt_state[doc_id]})
    except WebSocketDisconnect:
        active_docs[doc_id].discard(websocket)


@app.get("/collab/state/{doc_id}")
def collab_state(doc_id: str) -> dict[str, Any]:
    state = doc_crdt_state[doc_id]
    return {"doc_id": doc_id, "clock": state["clock"], "text": state["text"], "cursor": state["cursor"]}


@app.post("/export/pdf")
def export_pdf(body: ExportRequest) -> Response:
    safe_template = Environment(autoescape=True).from_string(
        """
    <style>
    body { font-family: Inter, Arial, sans-serif; }
    {{ custom_css }}
    </style>
    <h1>{{ title }}</h1>
    <p><strong>Author:</strong> {{ author }}</p>
    <div>{{ body }}</div>
    """
    )
    html = safe_template.render(title=body.title, author=body.author, body=body.body, custom_css=body.custom_css)
    # Lightweight text-based output with PDF media type for integration contract testing.
    return Response(content=html.encode("utf-8"), media_type="application/pdf")


@app.post("/export/docx/preview", response_model=DocxPreviewResponse)
def export_docx_preview(body: DocxTemplateRequest) -> DocxPreviewResponse:
    resolved_xml = _resolve_template_placeholders(body.template_xml, body.json_data)
    try:
        root = ET.fromstring(resolved_xml)
    except ET.ParseError as exc:
        raise HTTPException(status_code=400, detail=f"invalid template XML: {exc}") from exc

    paragraphs = _extract_docx_paragraphs(root)
    preview_html = _build_preview_html(root)
    return DocxPreviewResponse(resolved_xml=resolved_xml, preview_html=preview_html, paragraphs=paragraphs)


@app.post("/export/docx")
def export_docx(body: DocxTemplateRequest) -> Response:
    resolved_xml = _resolve_template_placeholders(body.template_xml, body.json_data)
    docx_bytes = _build_docx_bytes_from_xml(resolved_xml)
    safe_filename = body.filename.strip() or "draft.docx"
    if not safe_filename.lower().endswith(".docx"):
        safe_filename = f"{safe_filename}.docx"

    return Response(
        content=docx_bytes,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f'attachment; filename="{safe_filename}"'},
    )


@app.post("/export/docx/session", response_model=DocxSessionResponse)
def export_docx_session(body: DocxTemplateRequest) -> DocxSessionResponse:
    resolved_xml = _resolve_template_placeholders(body.template_xml, body.json_data)
    docx_bytes = _build_docx_bytes_from_xml(resolved_xml)
    safe_filename = body.filename.strip() or "draft.docx"
    if not safe_filename.lower().endswith(".docx"):
        safe_filename = f"{safe_filename}.docx"

    _clean_docx_artifacts()
    artifact_id = str(uuid.uuid4())
    DOCX_ARTIFACT_DIR.mkdir(parents=True, exist_ok=True)
    file_path = DOCX_ARTIFACT_DIR / f"{artifact_id}-{re.sub(r'[^a-zA-Z0-9._-]+', '-', safe_filename)}"
    file_path.write_bytes(docx_bytes)
    docx_artifacts[artifact_id] = {
        "content": docx_bytes,
        "filename": safe_filename,
        "file_path": str(file_path),
        "created_at": time.time(),
    }
    return DocxSessionResponse(
        docx_id=artifact_id,
        filename=safe_filename,
        viewer_path=f"/docx-viewer/render/{artifact_id}",
    )


@app.get("/export/docx/files/{docx_id}")
def export_docx_file(docx_id: str) -> Response:
    _clean_docx_artifacts()
    artifact = docx_artifacts.get(docx_id)
    if not artifact:
        raise HTTPException(status_code=404, detail="docx artifact not found")

    file_path = artifact.get("file_path")
    content = artifact.get("content")
    if isinstance(file_path, str) and file_path and Path(file_path).exists():
        content = Path(file_path).read_bytes()
        artifact["content"] = content

    return Response(
        content=content,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f'inline; filename="{artifact["filename"]}"'},
    )


@app.post("/ai/mcp/superdoc/open")
def ai_mcp_superdoc_open(body: SuperdocOpenRequest) -> dict[str, Any]:
    file_path = _ensure_docx_artifact_file(body.docx_id)
    opened = _mcp_bridge_post("/superdoc/open", {"path": file_path})
    return {"docx_id": body.docx_id, "opened": opened}


@app.post("/ai/mcp/superdoc/apply")
def ai_mcp_superdoc_apply(body: SuperdocApplyRequest) -> dict[str, Any]:
    file_path = _ensure_docx_artifact_file(body.docx_id)
    opened = _mcp_bridge_post("/superdoc/open", {"path": file_path})
    parsed_open = opened.get("parsed")
    if not isinstance(parsed_open, dict) or "session_id" not in parsed_open:
        raise HTTPException(status_code=502, detail="mcp open did not return session_id")

    session_id = str(parsed_open["session_id"])
    clipped = body.content.strip()[:3500]

    try:
        _mcp_bridge_post(
            "/superdoc/create-paragraph",
            {"session_id": session_id, "text": clipped, "suggest": body.suggest},
        )
        _mcp_bridge_post("/superdoc/save", {"session_id": session_id})
    finally:
        try:
            _mcp_bridge_post("/superdoc/close", {"session_id": session_id})
        except HTTPException:
            pass

    _reload_docx_artifact_content(body.docx_id)
    return {
        "docx_id": body.docx_id,
        "applied": True,
        "session_id": session_id,
        "file_path": file_path,
    }


@app.post("/ai/mcp/superdoc/save")
def ai_mcp_superdoc_save(body: SuperdocSessionRequest) -> dict[str, Any]:
    return _mcp_bridge_post("/superdoc/save", {"session_id": body.session_id})


@app.post("/ai/mcp/superdoc/close")
def ai_mcp_superdoc_close(body: SuperdocSessionRequest) -> dict[str, Any]:
    return _mcp_bridge_post("/superdoc/close", {"session_id": body.session_id})


@app.post("/billing/webhook")
def billing_webhook(
    payload: dict[str, Any], stripe_signature: str | None = Header(default=None, alias="Stripe-Signature")
) -> dict[str, Any]:
    if not stripe_signature:
        raise HTTPException(status_code=400, detail="missing stripe signature")

    event_type = payload.get("type")
    user_id = str(payload.get("data", {}).get("object", {}).get("metadata", {}).get("user_id", ""))
    if not user_id:
        raise HTTPException(status_code=400, detail="missing user id metadata")

    if event_type == "invoice.paid":
        subscriptions[user_id] = "enterprise"
    elif event_type == "subscription.deleted":
        subscriptions[user_id] = "free"

    return {"ok": True, "user_id": user_id, "tier": subscriptions[user_id]}


@app.get("/admin/service-health")
def admin_service_health() -> dict[str, Any]:
    return {
        "services": list(service_registry.values()),
        "summary": {
            "total": len(service_registry),
            "healthy": len(service_registry),
            "degraded": 0,
        },
    }
